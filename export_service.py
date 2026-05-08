import html
from io import BytesIO

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt


def normalize_saved_findings(saved):
    if isinstance(saved, dict):
        return {
            "format": saved.get("format") or "gov",
            "summary": saved.get("summary") or "",
            "items": saved.get("items") if isinstance(saved.get("items"), list) else [],
        }
    if isinstance(saved, list):
        return {"format": "local", "summary": "", "items": saved}
    return {"format": "gov", "summary": "", "items": []}


def build_findings_docx(session: dict) -> BytesIO:
    report = normalize_saved_findings(session.get("findings"))
    if not report["items"]:
        raise ValueError("No findings to export")

    document = Document()
    _set_docx_default_font(document)

    document.add_heading("資安稽核發現報告", level=0)
    _add_docx_meta(document, session, report)

    if report["summary"]:
        document.add_heading("摘要", level=1)
        document.add_paragraph(report["summary"])

    document.add_heading("稽核發現", level=1)
    for index, finding in enumerate(report["items"], start=1):
        _add_docx_finding(document, index, finding, report["format"])

    buffer = BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer


def build_findings_pdf(session: dict) -> BytesIO:
    report = normalize_saved_findings(session.get("findings"))
    if not report["items"]:
        raise ValueError("No findings to export")

    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import cm
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.cidfonts import UnicodeCIDFont
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    except ImportError as exc:
        raise RuntimeError("PDF export requires reportlab") from exc

    font_name = _register_cjk_font(pdfmetrics, UnicodeCIDFont)
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle("CJKTitle", parent=styles["Title"], fontName=font_name, fontSize=18, leading=24))
    styles.add(ParagraphStyle("CJKHeading", parent=styles["Heading2"], fontName=font_name, fontSize=13, leading=18, spaceBefore=10, spaceAfter=6))
    styles.add(ParagraphStyle("CJKBody", parent=styles["BodyText"], fontName=font_name, fontSize=10.5, leading=16))
    styles.add(ParagraphStyle("CJKSmall", parent=styles["BodyText"], fontName=font_name, fontSize=9, leading=13, textColor=colors.HexColor("#4b5563")))

    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=1.6 * cm,
        leftMargin=1.6 * cm,
        topMargin=1.5 * cm,
        bottomMargin=1.5 * cm,
        title="資安稽核發現報告",
    )
    story = [Paragraph("資安稽核發現報告", styles["CJKTitle"]), Spacer(1, 8)]
    story.extend(_pdf_meta_table(session, report, Paragraph, Table, TableStyle, styles, colors))

    if report["summary"]:
        story.extend([
            Spacer(1, 8),
            Paragraph("摘要", styles["CJKHeading"]),
            Paragraph(_pdf_escape(report["summary"]), styles["CJKBody"]),
        ])

    story.append(Spacer(1, 8))
    story.append(Paragraph("稽核發現", styles["CJKHeading"]))
    for index, finding in enumerate(report["items"], start=1):
        story.extend(_pdf_finding(index, finding, report["format"], Paragraph, Table, TableStyle, styles, colors))
        story.append(Spacer(1, 8))

    doc.build(story)
    buffer.seek(0)
    return buffer


def _set_docx_default_font(document: Document):
    styles = document.styles
    styles["Normal"].font.name = "Microsoft JhengHei"
    styles["Normal"].font.size = Pt(11)
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    for style_name in ("Heading 1", "Heading 2", "Heading 3"):
        style = styles[style_name]
        style.font.name = "Microsoft JhengHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")


def _add_docx_meta(document: Document, session: dict, report: dict):
    table = document.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    pairs = [
        ("紀錄編號", session.get("session_id", "")),
        ("稽核員", session.get("user_name", "")),
        ("報告格式", _format_label(report.get("format"))),
        ("稽核範圍", session.get("scope", "")),
        ("稽核情境", session.get("context", "")),
    ]
    for label, value in pairs:
        row = table.add_row().cells
        row[0].text = label
        row[1].text = str(value or "")
    document.add_paragraph("")


def _add_docx_finding(document: Document, index: int, finding: dict, report_format: str):
    title = finding.get("title") or finding.get("finding_type") or f"發現 {index}"
    document.add_heading(f"F{index} {title}", level=2)
    if report_format == "gov" and (finding.get("finding_type") or finding.get("finding_description")):
        fields = [
            ("類型", finding.get("finding_type")),
            ("法源依據", finding.get("legal_basis")),
            ("應辦事項（法條原文）", finding.get("legal_text")),
            ("稽核發現說明", finding.get("finding_description")),
            ("改善建議", finding.get("recommendation")),
        ]
        for label, value in fields:
            _add_docx_labeled_paragraph(document, label, value)
        evidence = finding.get("evidence") if isinstance(finding.get("evidence"), list) else []
        if evidence:
            document.add_paragraph("佐證資料：")
            for item in evidence:
                document.add_paragraph(str(item), style="List Bullet")
        return

    fields = [
        ("風險等級", finding.get("level") or finding.get("risk_level")),
        ("現況", finding.get("condition")),
        ("準則", finding.get("criteria") or finding.get("legal_basis")),
        ("原因", finding.get("cause")),
        ("影響", finding.get("effect")),
        ("改善建議", finding.get("recommendation")),
    ]
    for label, value in fields:
        _add_docx_labeled_paragraph(document, label, value)


def _add_docx_labeled_paragraph(document: Document, label: str, value):
    if not value:
        return
    paragraph = document.add_paragraph()
    paragraph.add_run(f"{label}：").bold = True
    paragraph.add_run(str(value))


def _pdf_meta_table(session, report, Paragraph, Table, TableStyle, styles, colors):
    rows = [
        ["紀錄編號", session.get("session_id", "")],
        ["稽核員", session.get("user_name", "")],
        ["報告格式", _format_label(report.get("format"))],
        ["稽核範圍", session.get("scope", "")],
        ["稽核情境", session.get("context", "")],
    ]
    table = Table(
        [[Paragraph(_pdf_escape(label), styles["CJKSmall"]), Paragraph(_pdf_escape(value), styles["CJKSmall"])] for label, value in rows],
        colWidths=[70, 390],
    )
    table.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#d1d5db")),
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#f3f4f6")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    return [table]


def _pdf_finding(index, finding, report_format, Paragraph, Table, TableStyle, styles, colors):
    title = finding.get("title") or finding.get("finding_type") or f"發現 {index}"
    story = [Paragraph(_pdf_escape(f"F{index} {title}"), styles["CJKHeading"])]
    if report_format == "gov" and (finding.get("finding_type") or finding.get("finding_description")):
        rows = [
            ("類型", finding.get("finding_type")),
            ("法源依據", finding.get("legal_basis")),
            ("應辦事項（法條原文）", finding.get("legal_text")),
            ("稽核發現說明", finding.get("finding_description")),
            ("佐證資料", "\n".join(str(item) for item in finding.get("evidence", []) if item) if isinstance(finding.get("evidence"), list) else ""),
            ("改善建議", finding.get("recommendation")),
        ]
    else:
        rows = [
            ("風險等級", finding.get("level") or finding.get("risk_level")),
            ("現況", finding.get("condition")),
            ("準則", finding.get("criteria") or finding.get("legal_basis")),
            ("原因", finding.get("cause")),
            ("影響", finding.get("effect")),
            ("改善建議", finding.get("recommendation")),
        ]
    filtered = [(label, value) for label, value in rows if value]
    table = Table(
        [[Paragraph(_pdf_escape(label), styles["CJKSmall"]), Paragraph(_pdf_escape(value), styles["CJKBody"])] for label, value in filtered],
        colWidths=[80, 380],
    )
    table.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#d1d5db")),
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#f9fafb")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    story.append(table)
    return story


def _register_cjk_font(pdfmetrics, UnicodeCIDFont):
    for name in ("MSung-Light", "STSong-Light"):
        try:
            pdfmetrics.registerFont(UnicodeCIDFont(name))
            return name
        except Exception:
            continue
    return "Helvetica"


def _pdf_escape(value) -> str:
    return html.escape(str(value or "")).replace("\n", "<br/>")


def _format_label(value) -> str:
    if value == "gov":
        return "衛福部/數位部 CI"
    if value == "iia5c":
        return "IIA 5C"
    if value == "local":
        return "規則草稿"
    return str(value or "")
